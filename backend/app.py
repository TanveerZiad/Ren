from __future__ import annotations

import hashlib
import html
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import threading
import urllib.error
import urllib.request
import uuid
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Literal

from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

from backend.chat_service import complete_chat
from backend.rag_service import (
    build_grounded_context_and_citations,
    chunk_source_document,
    compute_content_hash,
    embed_chunks_for_source,
    retrieve_hybrid_chunks,
)

APP_NAME = "Ren"
DEFAULT_DATA_DIR = Path(os.getenv("APPDATA", Path.home() / ".ren")) / "Ren"
DATA_DIR = Path(os.getenv("REN_DATA_DIR", DEFAULT_DATA_DIR))
LIBRARY_DIR = DATA_DIR / "library"
NOTES_DIR = DATA_DIR / "notes"
DB_PATH = DATA_DIR / "ren.sqlite3"


def now() -> str:
    return datetime.now(UTC).isoformat()


def ensure_directories() -> None:
    for directory in (DATA_DIR, LIBRARY_DIR, NOTES_DIR):
        directory.mkdir(parents=True, exist_ok=True)


@contextmanager
def db_connection():
    ensure_directories()
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    finally:
        connection.close()


def initialize_database() -> None:
    with db_connection() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS sources (
              id TEXT PRIMARY KEY, filename TEXT NOT NULL, source_kind TEXT NOT NULL,
              original_path TEXT, managed_path TEXT, url TEXT, mime_type TEXT,
              content_hash TEXT, extracted_text TEXT DEFAULT '', status TEXT NOT NULL,
              error_message TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE UNIQUE INDEX IF NOT EXISTS idx_sources_hash ON sources(content_hash)
              WHERE content_hash IS NOT NULL;
            CREATE TABLE IF NOT EXISTS processing_jobs (
              id TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES sources(id) ON DELETE CASCADE,
              status TEXT NOT NULL, stage TEXT NOT NULL, error_message TEXT,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS chunks (
              id TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES sources(id) ON DELETE CASCADE,
              ordinal INTEGER NOT NULL, text TEXT NOT NULL, anchor TEXT, content_hash TEXT, created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_chunks_source ON chunks(source_id);
            CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(chunk_id UNINDEXED, text);
            CREATE TABLE IF NOT EXISTS embeddings (
              id TEXT PRIMARY KEY,
              chunk_id TEXT NOT NULL UNIQUE REFERENCES chunks(id) ON DELETE CASCADE,
              model TEXT NOT NULL,
              dimensions INTEGER NOT NULL,
              vector_json TEXT NOT NULL,
              content_hash TEXT NOT NULL,
              created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_embeddings_hash ON embeddings(content_hash);
            CREATE TABLE IF NOT EXISTS topics (
              id TEXT PRIMARY KEY, name TEXT NOT NULL COLLATE NOCASE UNIQUE,
              summary TEXT DEFAULT '', created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS source_topics (
              source_id TEXT NOT NULL REFERENCES sources(id) ON DELETE CASCADE,
              topic_id TEXT NOT NULL REFERENCES topics(id) ON DELETE CASCADE,
              confidence REAL NOT NULL, PRIMARY KEY(source_id, topic_id)
            );
            CREATE TABLE IF NOT EXISTS notes (
              id TEXT PRIMARY KEY, topic_id TEXT NOT NULL UNIQUE REFERENCES topics(id) ON DELETE CASCADE,
              status TEXT NOT NULL, current_revision_id TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS note_revisions (
              id TEXT PRIMARY KEY, note_id TEXT NOT NULL REFERENCES notes(id) ON DELETE CASCADE,
              status TEXT NOT NULL, body TEXT NOT NULL, diff_summary TEXT NOT NULL,
              created_at TEXT NOT NULL, reviewed_at TEXT
            );
            CREATE TABLE IF NOT EXISTS citations (
              id TEXT PRIMARY KEY, revision_id TEXT REFERENCES note_revisions(id) ON DELETE CASCADE,
              source_id TEXT NOT NULL REFERENCES sources(id) ON DELETE CASCADE,
              chunk_id TEXT REFERENCES chunks(id) ON DELETE SET NULL, label TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS questions (
              id TEXT PRIMARY KEY, question TEXT NOT NULL, answer TEXT NOT NULL,
              created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS settings (
              key TEXT PRIMARY KEY, value TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS cards (
              id TEXT PRIMARY KEY,
              source_id TEXT NOT NULL REFERENCES sources(id) ON DELETE CASCADE,
              topic_id TEXT REFERENCES topics(id) ON DELETE SET NULL,
              title TEXT NOT NULL,
              body TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT 'active',
              last_seen_at TEXT,
              review_count INTEGER NOT NULL DEFAULT 0,
              next_review_at TEXT NOT NULL,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              UNIQUE(source_id, title)
            );
            CREATE INDEX IF NOT EXISTS idx_cards_review ON cards(status, next_review_at);
            CREATE TABLE IF NOT EXISTS folders (
              id TEXT PRIMARY KEY, name TEXT NOT NULL, position INTEGER NOT NULL DEFAULT 0,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS chat_threads (
              id TEXT PRIMARY KEY, title TEXT NOT NULL, folder_id TEXT REFERENCES folders(id) ON DELETE SET NULL,
              incognito INTEGER NOT NULL DEFAULT 0, locked INTEGER NOT NULL DEFAULT 0,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_chat_threads_updated ON chat_threads(updated_at DESC);
            CREATE TABLE IF NOT EXISTS chat_messages (
              id TEXT PRIMARY KEY, thread_id TEXT NOT NULL REFERENCES chat_threads(id) ON DELETE CASCADE,
              role TEXT NOT NULL CHECK(role IN ('user', 'assistant', 'system')),
              content TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'complete', model TEXT,
              citations_json TEXT NOT NULL DEFAULT '[]', created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_chat_messages_thread ON chat_messages(thread_id, created_at);
            CREATE TABLE IF NOT EXISTS memories (
              id TEXT PRIMARY KEY, content TEXT NOT NULL, category TEXT NOT NULL DEFAULT 'general',
              source_thread_id TEXT REFERENCES chat_threads(id) ON DELETE SET NULL,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS artifacts (
              id TEXT PRIMARY KEY, thread_id TEXT NOT NULL REFERENCES chat_threads(id) ON DELETE CASCADE,
              title TEXT NOT NULL, kind TEXT NOT NULL CHECK(kind IN ('document', 'html', 'chart')),
              content TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'complete', share_id TEXT UNIQUE,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS integrations (
              id TEXT PRIMARY KEY, name TEXT NOT NULL, endpoint TEXT NOT NULL,
              auth_mode TEXT NOT NULL DEFAULT 'none', enabled INTEGER NOT NULL DEFAULT 0,
              status TEXT NOT NULL DEFAULT 'not_connected', created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS queued_messages (
              id TEXT PRIMARY KEY, thread_id TEXT NOT NULL REFERENCES chat_threads(id) ON DELETE CASCADE,
              content TEXT NOT NULL, created_at TEXT NOT NULL
            );
            """
        )
        try:
            db.execute("ALTER TABLE chunks ADD COLUMN content_hash TEXT")
        except sqlite3.OperationalError:
            pass
        try:
            db.execute("CREATE INDEX IF NOT EXISTS idx_chunks_hash ON chunks(content_hash)")
        except sqlite3.OperationalError:
            pass


class StripHtml(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self.skip_depth += 1
        elif tag in {"p", "br", "div", "h1", "h2", "h3", "li", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self.skip_depth:
            self.skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self.skip_depth:
            self.parts.append(data)


def extract_text(path: Path, source_kind: str) -> str:
    suffix = path.suffix.lower()
    if source_kind == "text" or suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        parser = StripHtml()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return html.unescape(" ".join(parser.parts))
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        except Exception as error:
            raise ValueError(f"PDF extraction failed: {error}") from error
    if suffix == ".docx":
        try:
            from docx import Document
            return "\n".join(p.text for p in Document(str(path)).paragraphs)
        except Exception as error:
            raise ValueError(f"DOCX extraction failed: {error}") from error
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def normalize_text(text: str) -> str:
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", text)).strip()


def make_chunks(text: str, size: int = 1100, overlap: int = 180) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = max(text.rfind(". ", start, end), text.rfind("\n", start, end))
            if boundary > start + size // 2:
                end = boundary + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [chunk for chunk in chunks if chunk]


STOP_WORDS = {"about", "after", "again", "also", "and", "are", "been", "from", "have", "into", "more", "notes", "that", "this", "with", "your", "the", "for", "how", "what"}
CARD_STOP_WORDS = STOP_WORDS | {
    "able", "been", "being", "between", "both", "can", "could", "does", "each", "first",
    "however", "into", "its", "may", "most", "must", "not", "only", "other", "over",
    "such", "than", "their", "there", "these", "they", "through", "using", "very", "were",
    "when", "which", "while", "will", "would", "you", "also", "just", "like",
}


def suggested_topics(filename: str, text: str) -> list[str]:
    candidates = re.findall(r"^#{1,3}\s+(.+)$", text, flags=re.MULTILINE)
    candidates += re.findall(r"\b[A-Z][A-Za-z0-9+\-]{2,}(?:\s+[A-Z][A-Za-z0-9+\-]{2,}){0,2}\b", text[:8000])
    candidates += re.split(r"[_\-\.]+", Path(filename).stem)
    unique: list[str] = []
    for candidate in candidates:
        candidate = re.sub(r"\s+", " ", candidate).strip(" -_:.")
        if not candidate or candidate.lower() in STOP_WORDS or len(candidate) < 3 or len(candidate) > 60:
            continue
        if candidate.lower() not in {item.lower() for item in unique}:
            unique.append(candidate)
        if len(unique) == 5:
            break
    return unique or [Path(filename).stem[:60] or "Inbox"]


def get_or_create_topic(db: sqlite3.Connection, name: str) -> str:
    row = db.execute("SELECT id FROM topics WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    topic_id = str(uuid.uuid4())
    db.execute("INSERT INTO topics(id, name, created_at) VALUES (?, ?, ?)", (topic_id, name, now()))
    return topic_id


def update_job(db: sqlite3.Connection, job_id: str, status: str, stage: str, error: str | None = None) -> None:
    db.execute("UPDATE processing_jobs SET status=?, stage=?, error_message=?, updated_at=? WHERE id=?", (status, stage, error, now(), job_id))


def configured_embedding_provider() -> tuple[str, str]:
    with db_connection() as db:
        values = {row["key"]: row["value"] for row in db.execute("SELECT key, value FROM settings")}
    return values.get("base_url", "https://openrouter.ai/api/v1"), values.get("embedding_model", "text-embedding-3-small")


def process_source(source_id: str, job_id: str) -> None:
    affected_topics: list[str] = []
    with db_connection() as db:
        source = db.execute("SELECT * FROM sources WHERE id=?", (source_id,)).fetchone()
        if not source:
            return
        try:
            update_job(db, job_id, "processing", "extracting")
            text = source["extracted_text"] or extract_text(Path(source["managed_path"]), source["source_kind"])
            text = normalize_text(text)
            if not text:
                raise ValueError("No readable text was found in this source.")
            db.execute("UPDATE sources SET extracted_text=?, status='processing', updated_at=? WHERE id=?", (text, now(), source_id))

            update_job(db, job_id, "processing", "chunking")
            managed_path = Path(source["managed_path"]) if source["managed_path"] else None
            chunks_data = chunk_source_document(managed_path, text, source["source_kind"])
            if not chunks_data:
                chunks_data = [{"ordinal": 1, "text": text, "anchor": "Chunk 1", "content_hash": compute_content_hash(text)}]

            base_url, embedding_model = configured_embedding_provider()
            all_hashes = list({c["content_hash"] for c in chunks_data})
            known_vectors: dict[str, tuple[int, str]] = {}
            if all_hashes:
                placeholders = ",".join("?" for _ in all_hashes)
                existing_rows = db.execute(
                    f"SELECT content_hash, dimensions, vector_json FROM embeddings WHERE content_hash IN ({placeholders}) AND model=?",
                    all_hashes + [embedding_model],
                ).fetchall()
                for r in existing_rows:
                    known_vectors[r["content_hash"]] = (r["dimensions"], r["vector_json"])

            db.execute("DELETE FROM chunks_fts WHERE chunk_id IN (SELECT id FROM chunks WHERE source_id=?)", (source_id,))
            db.execute("DELETE FROM chunks WHERE source_id=?", (source_id,))

            created_chunks = []
            for c in chunks_data:
                chunk_id = str(uuid.uuid4())
                db.execute(
                    "INSERT INTO chunks(id, source_id, ordinal, text, anchor, content_hash, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                    (chunk_id, source_id, c["ordinal"], c["text"], c["anchor"], c["content_hash"], now()),
                )
                db.execute("INSERT INTO chunks_fts(chunk_id, text) VALUES (?, ?)", (chunk_id, c["text"]))
                created_chunks.append({
                    "id": chunk_id,
                    "chunk_id": chunk_id,
                    "source_id": source_id,
                    "ordinal": c["ordinal"],
                    "text": c["text"],
                    "anchor": c["anchor"],
                    "content_hash": c["content_hash"],
                })

            update_job(db, job_id, "processing", "embedding")
            newly, reused, embed_error = embed_chunks_for_source(
                db=db,
                chunks=created_chunks,
                known_vectors=known_vectors,
                base_url=base_url,
                api_key=get_api_key(),
                model=embedding_model,
            )

            update_job(db, job_id, "processing", "suggesting topics")
            for topic_name in suggested_topics(source["filename"], text):
                topic_id = get_or_create_topic(db, topic_name)
                affected_topics.append(topic_id)
                db.execute("INSERT OR REPLACE INTO source_topics(source_id, topic_id, confidence) VALUES (?, ?, ?)", (source_id, topic_id, 0.65))
            db.execute("UPDATE sources SET status='ready', error_message=NULL, updated_at=? WHERE id=?", (now(), source_id))
            if embed_error:
                update_job(db, job_id, "completed", "ready", embed_error)
            else:
                update_job(db, job_id, "completed", "ready")
        except Exception as error:
            db.execute("UPDATE sources SET status='failed', error_message=?, updated_at=? WHERE id=?", (str(error), now(), source_id))
            update_job(db, job_id, "failed", "failed", str(error))
    # Ren's small-MVP output is reusable cards, not a large generated document.
    # This happens after indexing so source writes do not compete in SQLite.
    if affected_topics:
        create_cards_for_source(source_id)


def source_sentences(text: str) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+|\n{2,}", normalize_text(text))
    useful = [re.sub(r"\s+", " ", sentence).strip() for sentence in sentences]
    return [sentence for sentence in useful if 35 <= len(sentence) <= 520]


def source_title(filename: str) -> str:
    title = re.sub(r"[_\-.]+", " ", Path(filename).stem).strip()
    return title[:60].title() or "Imported knowledge"


def markdown_sections(text: str) -> list[tuple[str, str]]:
    """Return compact cards from human-authored Markdown headings when present."""
    matches = list(re.finditer(r"^#{1,4}\s+(.+?)\s*$", text, flags=re.MULTILINE))
    sections: list[tuple[str, str]] = []
    for index, match in enumerate(matches[:5]):
        title = re.sub(r"\s+", " ", match.group(1)).strip(" #:-")
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = normalize_text(text[match.end():end])
        if 3 <= len(title) <= 80 and len(body) >= 35:
            sections.append((title, body[:700]))
    return sections


def keyword_counts(text: str) -> dict[str, int]:
    words = re.findall(r"[A-Za-z][A-Za-z0-9+\-]{2,}", text.lower())
    counts: dict[str, int] = {}
    for word in words:
        if word not in CARD_STOP_WORDS:
            counts[word] = counts.get(word, 0) + 1
    return counts


def sentence_score(sentence: str, counts: dict[str, int]) -> float:
    words = set(re.findall(r"[A-Za-z][A-Za-z0-9+\-]{2,}", sentence.lower()))
    score = sum(min(counts.get(word, 0), 4) for word in words if word not in CARD_STOP_WORDS)
    # Definitions and explicit statements are usually more useful to revisit.
    if re.search(r"\b(is|are|means|refers to|uses|helps|enables|because)\b", sentence, re.I):
        score += 4
    if ":" in sentence:
        score += 1
    return score / max(len(words), 1)


def compact_card_title(sentence: str, counts: dict[str, int], number: int) -> str:
    words = re.findall(r"[A-Za-z][A-Za-z0-9+\-]{2,}", sentence)
    ranked = sorted(
        {word.lower() for word in words if word.lower() not in CARD_STOP_WORDS},
        key=lambda word: (-counts.get(word, 0), word),
    )
    if ranked:
        return " · ".join(word.title() for word in ranked[:2])
    return f"Key idea {number}"


def card_candidates(filename: str, text: str) -> list[tuple[str, str]]:
    sections = markdown_sections(text)
    if sections:
        return sections[:5]
    sentences = source_sentences(text)
    counts = keyword_counts(text)
    scored = sorted(enumerate(sentences), key=lambda item: sentence_score(item[1], counts), reverse=True)
    selected: list[tuple[str, str]] = []
    used_words: set[str] = set()
    for _, sentence in scored:
        words = set(re.findall(r"[A-Za-z][A-Za-z0-9+\-]{2,}", sentence.lower()))
        # Avoid five cards that merely repeat the same idea.
        if used_words and len(words & used_words) / max(len(words), 1) > 0.72:
            continue
        selected.append((compact_card_title(sentence, counts, len(selected) + 1), sentence[:700]))
        used_words |= words
        if len(selected) == 5:
            break
    if selected:
        return selected
    fallback = normalize_text(text)[:700]
    return [(source_title(filename), fallback)] if fallback else []


def create_cards_for_source(source_id: str) -> None:
    with db_connection() as db:
        source = db.execute("SELECT * FROM sources WHERE id=? AND status='ready'", (source_id,)).fetchone()
        if not source:
            return
        db.execute("DELETE FROM cards WHERE source_id=?", (source_id,))
        timestamp = now()
        topic_id = get_or_create_topic(db, source_title(source["filename"]))
        for title, body in card_candidates(source["filename"], source["extracted_text"]):
            db.execute(
                """INSERT INTO cards(id, source_id, topic_id, title, body, status, next_review_at, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, 'active', ?, ?, ?)""",
                (
                    str(uuid.uuid4()), source_id, topic_id, title, body,
                    timestamp, timestamp, timestamp,
                ),
            )


def card_payload(row: sqlite3.Row) -> dict[str, Any]:
    payload = dict(row)
    payload["source"] = {"id": row["source_id"], "filename": row["filename"]}
    return payload


def review_timestamp(review_count: int) -> str:
    intervals = (1, 3, 7, 14, 30)
    days = intervals[min(max(review_count - 1, 0), len(intervals) - 1)]
    return (datetime.now(UTC) + timedelta(days=days)).isoformat()


def query_rows(query: str, limit: int = 6) -> list[dict[str, Any]]:
    base_url, embedding_model = configured_embedding_provider()
    with db_connection() as db:
        return retrieve_hybrid_chunks(
            db=db,
            query=query,
            base_url=base_url,
            api_key=get_api_key(),
            embedding_model=embedding_model,
            limit=limit,
        )


def create_note_draft(topic_id: str) -> dict[str, Any]:
    with db_connection() as db:
        topic = db.execute("SELECT * FROM topics WHERE id=?", (topic_id,)).fetchone()
        if not topic:
            raise KeyError(topic_id)
        pending = db.execute(
            """SELECT r.*, n.id AS resolved_note_id FROM notes n JOIN note_revisions r ON r.note_id=n.id
               WHERE n.topic_id=? AND r.status='draft' ORDER BY r.created_at DESC LIMIT 1""", (topic_id,)
        ).fetchone()
        if pending:
            return {"id": pending["id"], "note_id": pending["resolved_note_id"], "topic": topic["name"], "body": pending["body"], "status": "draft", "diff_summary": pending["diff_summary"]}
        sources = db.execute(
            """SELECT s.* FROM sources s JOIN source_topics st ON st.source_id=s.id
               WHERE st.topic_id=? AND s.status='ready' ORDER BY s.created_at DESC""", (topic_id,)
        ).fetchall()
        if not sources:
            raise ValueError("This topic has no processed sources yet.")
        facts = []
        for source in sources[:6]:
            excerpt = normalize_text(source["extracted_text"])[:420]
            facts.append(f"- {excerpt}")
        source_lines = [f"- [{source['filename']}](source:{source['id']})" for source in sources]
        body = f"""# {topic['name']}\n\n## Overview\nThis living note is based on {len(sources)} source(s) currently linked to **{topic['name']}**. Review and edit this draft before approval.\n\n## Current understanding\n{chr(10).join(facts)}\n\n## Open questions\n- What should I learn or verify next about {topic['name']}?\n\n## Sources\n{chr(10).join(source_lines)}\n\n_Last generated by Ren on {datetime.now().strftime('%B %d, %Y')}_\n"""
        note = db.execute("SELECT * FROM notes WHERE topic_id=?", (topic_id,)).fetchone()
        if not note:
            note_id = str(uuid.uuid4())
            db.execute("INSERT INTO notes(id, topic_id, status, created_at, updated_at) VALUES (?, ?, 'draft', ?, ?)", (note_id, topic_id, now(), now()))
        else:
            note_id = note["id"]
        revision_id = str(uuid.uuid4())
        db.execute("INSERT INTO note_revisions(id, note_id, status, body, diff_summary, created_at) VALUES (?, ?, 'draft', ?, ?, ?)", (revision_id, note_id, body, "Generated from newly linked sources.", now()))
        for source in sources:
            db.execute("INSERT INTO citations(id, revision_id, source_id, label) VALUES (?, ?, ?, ?)", (str(uuid.uuid4()), revision_id, source["id"], source["filename"]))
        db.execute("UPDATE notes SET status='draft', updated_at=? WHERE id=?", (now(), note_id))
        return {"id": revision_id, "note_id": note_id, "topic": topic["name"], "body": body, "status": "draft", "diff_summary": "Generated from newly linked sources."}


class TextCapture(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    content: str = Field(min_length=1)


class UrlCapture(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    url: str = Field(min_length=3, max_length=2048)
    context: str = ""


class SettingsInput(BaseModel):
    provider: str = "openrouter"
    base_url: str = "https://openrouter.ai/api/v1"
    model: str = "openai/gpt-4o-mini"
    embedding_model: str = "text-embedding-3-small"
    api_key: str | None = None


class QuestionInput(BaseModel):
    question: str = Field(min_length=2, max_length=2000)


class ChatThreadInput(BaseModel):
    title: str = Field(default="New chat", min_length=1, max_length=120)
    incognito: bool = False
    folder_id: str | None = None


class ChatMessageInput(BaseModel):
    content: str = Field(min_length=1, max_length=12000)
    thinking: Literal["focused", "balanced", "creative"] = "balanced"
    model: str | None = Field(default=None, max_length=160)


class MemoryInput(BaseModel):
    content: str = Field(min_length=1, max_length=2000)
    category: str = Field(default="general", max_length=40)
    source_thread_id: str | None = None


class ArtifactInput(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    kind: Literal["document", "html", "chart"] = "document"
    content: str = Field(min_length=1, max_length=100000)


class FolderInput(BaseModel):
    name: str = Field(min_length=1, max_length=80)


class IntegrationInput(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    endpoint: str = Field(min_length=3, max_length=2048)
    auth_mode: Literal["none", "api_key", "oauth"] = "none"


app = FastAPI(title="Ren Local Service", version="0.1.0")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173", "tauri://localhost"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])
initialize_database()
_session_api_key: str | None = None
_keyring_unavailable: bool = False


def store_api_key(value: str) -> bool:
    global _session_api_key, _keyring_unavailable
    _keyring_unavailable = False
    try:
        import keyring
        keyring.set_password("Ren", "ai_api_key", value)
        _session_api_key = None
        return True
    except Exception:
        _session_api_key = value
        _keyring_unavailable = True
        return False


def api_key_available() -> bool:
    global _keyring_unavailable
    if _session_api_key:
        return True
    if _keyring_unavailable:
        return False
    try:
        import keyring
        return bool(keyring.get_password("Ren", "ai_api_key"))
    except Exception:
        _keyring_unavailable = True
        return False


def get_api_key() -> str | None:
    global _keyring_unavailable
    if _session_api_key:
        return _session_api_key
    if _keyring_unavailable:
        return None
    try:
        import keyring
        return keyring.get_password("Ren", "ai_api_key")
    except Exception:
        _keyring_unavailable = True
        return None


def provider_answer(question: str, rows: list[dict[str, Any]]) -> str | None:
    """Ask the configured provider using only already-retrieved local context.

    Returning None deliberately selects the source-extractive fallback. A provider
    failure must never prevent a user from retrieving their own information.
    """
    key = get_api_key()
    if not key:
        return None
    with db_connection() as db:
        configured = {row["key"]: row["value"] for row in db.execute("SELECT key, value FROM settings")}
    context, _ = build_grounded_context_and_citations(rows[:6])
    prompt = (
        "Answer only from the provided Ren source excerpts. "
        "If the evidence is insufficient, say so. "
        "Cite every important claim as [S1], [S2], etc. "
        "Do not use outside knowledge unless web search is explicitly enabled.\n\n"
        f"Question: {question}\n\nSource excerpts:\n{context}"
    )
    payload = json.dumps({
        "model": configured.get("model", "openai/gpt-4o-mini"),
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.2,
    }).encode("utf-8")
    base_url = configured.get("base_url", "https://openrouter.ai/api/v1").rstrip("/")
    request = urllib.request.Request(
        f"{base_url}/chat/completions", data=payload,
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"}, method="POST"
    )
    try:
        with urllib.request.urlopen(request, timeout=40) as response:
            data = json.loads(response.read().decode("utf-8"))
        return data["choices"][0]["message"]["content"].strip()
    except (urllib.error.URLError, urllib.error.HTTPError, KeyError, IndexError, json.JSONDecodeError):
        return None


def configured_provider() -> tuple[str, str]:
    with db_connection() as db:
        values = {row["key"]: row["value"] for row in db.execute("SELECT key, value FROM settings")}
    # OpenRouter accepts this API shape; the setting remains compatible with a
    # different OpenAI-compatible provider for a self-hosted future.
    return values.get("base_url", "https://openrouter.ai/api/v1"), values.get("model", "openai/gpt-4o-mini")


def chat_thread_payload(row: sqlite3.Row) -> dict[str, Any]:
    payload = dict(row)
    payload["incognito"] = bool(payload["incognito"])
    payload["locked"] = bool(payload["locked"])
    return payload


def chat_message_payload(row: sqlite3.Row) -> dict[str, Any]:
    payload = dict(row)
    payload["citations"] = json.loads(payload.pop("citations_json", "[]"))
    return payload


def chat_context(query: str) -> tuple[str, list[dict[str, Any]]]:
    rows = query_rows(query, limit=6)
    return build_grounded_context_and_citations(rows)


def source_payload(row: sqlite3.Row) -> dict[str, Any]:
    with db_connection() as db:
        card_count = db.execute("SELECT COUNT(*) AS total FROM cards WHERE source_id=? AND status='active'", (row["id"],)).fetchone()["total"]
    return dict(row) | {
        "topics": [item["name"] for item in get_source_topics(row["id"])],
        "card_count": card_count,
    }


def get_source_topics(source_id: str) -> list[sqlite3.Row]:
    with db_connection() as db:
        return db.execute("SELECT t.* FROM topics t JOIN source_topics st ON st.topic_id=t.id WHERE st.source_id=?", (source_id,)).fetchall()


def create_source(filename: str, source_kind: str, content: bytes, original_path: str | None = None, url: str | None = None) -> tuple[str, str | None]:
    digest = hashlib.sha256(content).hexdigest()
    with db_connection() as db:
        duplicate = db.execute("SELECT id FROM sources WHERE content_hash=?", (digest,)).fetchone()
        if duplicate:
            return duplicate["id"], "duplicate"
        source_id = str(uuid.uuid4())
        suffix = Path(filename).suffix
        managed_path = LIBRARY_DIR / f"{source_id}{suffix}"
        managed_path.write_bytes(content)
        timestamp = now()
        db.execute("""INSERT INTO sources(id, filename, source_kind, original_path, managed_path, url, mime_type, content_hash, status, created_at, updated_at)
                      VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'queued', ?, ?)""", (source_id, filename, source_kind, original_path, str(managed_path), url, mimetypes.guess_type(filename)[0], digest, timestamp, timestamp))
        return source_id, None


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"name": APP_NAME, "status": "ok", "data_dir": str(DATA_DIR)}


@app.get("/api/settings")
def get_settings() -> dict[str, Any]:
    with db_connection() as db:
        values = {row["key"]: row["value"] for row in db.execute("SELECT key, value FROM settings")}
    return {
        "provider": values.get("provider", "openrouter"),
        "base_url": values.get("base_url", "https://openrouter.ai/api/v1"),
        "model": values.get("model", "openai/gpt-4o-mini"),
        "embedding_model": values.get("embedding_model", "text-embedding-3-small"),
        "api_key_configured": api_key_available(),
    }


@app.put("/api/settings")
def save_settings(payload: SettingsInput) -> dict[str, Any]:
    with db_connection() as db:
        for key, value in {
            "provider": payload.provider.strip() or "openrouter",
            "base_url": payload.base_url.rstrip("/"),
            "model": payload.model.strip(),
            "embedding_model": (payload.embedding_model or "text-embedding-3-small").strip(),
        }.items():
            db.execute("INSERT INTO settings(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))
    vault_saved = None
    if payload.api_key:
        vault_saved = store_api_key(payload.api_key)
    return {
        "provider": payload.provider.strip() or "openrouter",
        "base_url": payload.base_url.rstrip("/"),
        "model": payload.model.strip(),
        "embedding_model": (payload.embedding_model or "text-embedding-3-small").strip(),
        "api_key_configured": api_key_available(),
        "vault_saved": vault_saved,
    }


@app.get("/api/platform/capabilities")
def platform_capabilities() -> dict[str, Any]:
    """Honest feature availability: no account or provider is assumed."""
    return {
        "chat": True,
        "knowledge": True,
        "memory": True,
        "artifacts": True,
        "folders": True,
        "incognito": True,
        "openrouter": api_key_available(),
        "mcp": False,
        "web_search": False,
        "voice": False,
        "image_generation": False,
        "sharing": False,
        "locked_chats": False,
    }


@app.get("/api/chat/threads")
def list_chat_threads() -> list[dict[str, Any]]:
    with db_connection() as db:
        rows = db.execute("SELECT * FROM chat_threads ORDER BY updated_at DESC").fetchall()
    return [chat_thread_payload(row) for row in rows]


@app.post("/api/chat/threads")
def create_chat_thread(payload: ChatThreadInput) -> dict[str, Any]:
    thread_id = str(uuid.uuid4())
    timestamp = now()
    with db_connection() as db:
        if payload.folder_id and not db.execute("SELECT 1 FROM folders WHERE id=?", (payload.folder_id,)).fetchone():
            raise HTTPException(422, "That folder no longer exists.")
        db.execute(
            """INSERT INTO chat_threads(id, title, folder_id, incognito, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (thread_id, payload.title.strip(), payload.folder_id, int(payload.incognito), timestamp, timestamp),
        )
        row = db.execute("SELECT * FROM chat_threads WHERE id=?", (thread_id,)).fetchone()
    return chat_thread_payload(row)


@app.get("/api/chat/threads/{thread_id}")
def get_chat_thread(thread_id: str) -> dict[str, Any]:
    with db_connection() as db:
        thread = db.execute("SELECT * FROM chat_threads WHERE id=?", (thread_id,)).fetchone()
        if not thread:
            raise HTTPException(404, "Chat not found")
        messages = db.execute("SELECT * FROM chat_messages WHERE thread_id=? ORDER BY created_at", (thread_id,)).fetchall()
    return {"thread": chat_thread_payload(thread), "messages": [chat_message_payload(row) for row in messages]}


@app.delete("/api/chat/threads/{thread_id}")
def delete_chat_thread(thread_id: str) -> dict[str, bool]:
    with db_connection() as db:
        deleted = db.execute("DELETE FROM chat_threads WHERE id=?", (thread_id,)).rowcount
    if not deleted:
        raise HTTPException(404, "Chat not found")
    return {"deleted": True}


@app.post("/api/chat/threads/{thread_id}/messages")
def send_chat_message(thread_id: str, payload: ChatMessageInput) -> dict[str, Any]:
    text = payload.content.strip()
    with db_connection() as db:
        thread = db.execute("SELECT * FROM chat_threads WHERE id=?", (thread_id,)).fetchone()
        if not thread:
            raise HTTPException(404, "Chat not found")
        if thread["locked"]:
            raise HTTPException(423, "Locked chats are not available in the local Ren build yet.")
        timestamp = now()
        user_id = str(uuid.uuid4())
        db.execute(
            "INSERT INTO chat_messages(id, thread_id, role, content, status, created_at) VALUES (?, ?, 'user', ?, 'complete', ?)",
            (user_id, thread_id, text, timestamp),
        )
        conversation_rows = db.execute("SELECT role, content FROM chat_messages WHERE thread_id=? ORDER BY created_at DESC LIMIT 16", (thread_id,)).fetchall()
        memories = [] if thread["incognito"] else [row["content"] for row in db.execute("SELECT content FROM memories ORDER BY updated_at DESC LIMIT 8")]
    context, citations = chat_context(text)
    base_url, default_model = configured_provider()
    completion = complete_chat(
        api_key=get_api_key(),
        base_url=base_url,
        model=payload.model or default_model,
        conversation=[{"role": row["role"], "content": row["content"]} for row in reversed(conversation_rows)],
        context=context,
        memories=memories,
        thinking=payload.thinking,
    )
    assistant_id = str(uuid.uuid4())
    with db_connection() as db:
        db.execute(
            """INSERT INTO chat_messages(id, thread_id, role, content, status, model, citations_json, created_at)
               VALUES (?, ?, 'assistant', ?, ?, ?, ?, ?)""",
            (assistant_id, thread_id, completion.content, "complete" if completion.error is None else "failed", payload.model or default_model, json.dumps(citations), now()),
        )
        db.execute("UPDATE chat_threads SET updated_at=? WHERE id=?", (now(), thread_id))
        assistant = db.execute("SELECT * FROM chat_messages WHERE id=?", (assistant_id,)).fetchone()
    return chat_message_payload(assistant) | {"provider_used": completion.provider_used, "error": completion.error}


@app.get("/api/memories")
def list_memories() -> list[dict[str, Any]]:
    with db_connection() as db:
        return [dict(row) for row in db.execute("SELECT * FROM memories ORDER BY updated_at DESC")]


@app.post("/api/memories")
def create_memory(payload: MemoryInput) -> dict[str, Any]:
    memory_id = str(uuid.uuid4())
    timestamp = now()
    with db_connection() as db:
        db.execute("INSERT INTO memories(id, content, category, source_thread_id, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)", (memory_id, payload.content.strip(), payload.category.strip(), payload.source_thread_id, timestamp, timestamp))
        row = db.execute("SELECT * FROM memories WHERE id=?", (memory_id,)).fetchone()
    return dict(row)


@app.delete("/api/memories/{memory_id}")
def delete_memory(memory_id: str) -> dict[str, bool]:
    with db_connection() as db:
        deleted = db.execute("DELETE FROM memories WHERE id=?", (memory_id,)).rowcount
    if not deleted:
        raise HTTPException(404, "Memory not found")
    return {"deleted": True}


@app.get("/api/chat/threads/{thread_id}/artifacts")
def list_artifacts(thread_id: str) -> list[dict[str, Any]]:
    with db_connection() as db:
        return [dict(row) for row in db.execute("SELECT * FROM artifacts WHERE thread_id=? ORDER BY updated_at DESC", (thread_id,))]


@app.post("/api/chat/threads/{thread_id}/artifacts")
def create_artifact(thread_id: str, payload: ArtifactInput) -> dict[str, Any]:
    artifact_id = str(uuid.uuid4())
    timestamp = now()
    with db_connection() as db:
        if not db.execute("SELECT 1 FROM chat_threads WHERE id=?", (thread_id,)).fetchone():
            raise HTTPException(404, "Chat not found")
        db.execute("INSERT INTO artifacts(id, thread_id, title, kind, content, share_id, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (artifact_id, thread_id, payload.title.strip(), payload.kind, payload.content, uuid.uuid4().hex[:10], timestamp, timestamp))
        row = db.execute("SELECT * FROM artifacts WHERE id=?", (artifact_id,)).fetchone()
    return dict(row)


@app.post("/api/folders")
def create_folder(payload: FolderInput) -> dict[str, Any]:
    folder_id = str(uuid.uuid4())
    timestamp = now()
    with db_connection() as db:
        position = db.execute("SELECT COUNT(*) AS total FROM folders").fetchone()["total"]
        db.execute("INSERT INTO folders(id, name, position, created_at, updated_at) VALUES (?, ?, ?, ?, ?)", (folder_id, payload.name.strip(), position, timestamp, timestamp))
        row = db.execute("SELECT * FROM folders WHERE id=?", (folder_id,)).fetchone()
    return dict(row)


@app.get("/api/folders")
def list_folders() -> list[dict[str, Any]]:
    with db_connection() as db:
        return [dict(row) for row in db.execute("SELECT * FROM folders ORDER BY position, name")]


@app.post("/api/integrations")
def register_integration(payload: IntegrationInput) -> dict[str, Any]:
    integration_id = str(uuid.uuid4())
    timestamp = now()
    with db_connection() as db:
        db.execute("INSERT INTO integrations(id, name, endpoint, auth_mode, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)", (integration_id, payload.name.strip(), payload.endpoint.strip(), payload.auth_mode, timestamp, timestamp))
        row = db.execute("SELECT * FROM integrations WHERE id=?", (integration_id,)).fetchone()
    return dict(row)


@app.get("/api/integrations")
def list_integrations() -> list[dict[str, Any]]:
    with db_connection() as db:
        return [dict(row) for row in db.execute("SELECT * FROM integrations ORDER BY updated_at DESC")]


@app.post("/api/captures/files")
async def import_files(background_tasks: BackgroundTasks, files: list[UploadFile] = File(...)) -> list[dict[str, Any]]:
    results = []
    allowed_extensions = {".pdf", ".txt", ".md", ".markdown"}
    for upload in files:
        filename = upload.filename or "Untitled"
        if Path(filename).suffix.lower() not in allowed_extensions:
            results.append({"filename": filename, "status": "unsupported", "error": "Ren currently supports PDF, TXT, and Markdown files."})
            continue
        content = await upload.read()
        source_id, duplicate = create_source(filename, "file", content)
        if duplicate:
            results.append({"id": source_id, "status": "duplicate"})
            continue
        job_id = str(uuid.uuid4())
        with db_connection() as db:
            db.execute("INSERT INTO processing_jobs(id, source_id, status, stage, created_at, updated_at) VALUES (?, ?, 'queued', 'waiting', ?, ?)", (job_id, source_id, now(), now()))
        background_tasks.add_task(process_source, source_id, job_id)
        results.append({"id": source_id, "job_id": job_id, "status": "queued"})
    return results


@app.post("/api/captures/text")
def import_text(payload: TextCapture, background_tasks: BackgroundTasks) -> dict[str, Any]:
    filename = f"{payload.title.strip()[:100]}.md"
    source_id, duplicate = create_source(filename, "text", payload.content.encode("utf-8"))
    if duplicate:
        return {"id": source_id, "status": "duplicate"}
    job_id = str(uuid.uuid4())
    with db_connection() as db:
        db.execute("INSERT INTO processing_jobs(id, source_id, status, stage, created_at, updated_at) VALUES (?, ?, 'queued', 'waiting', ?, ?)", (job_id, source_id, now(), now()))
    background_tasks.add_task(process_source, source_id, job_id)
    return {"id": source_id, "job_id": job_id, "status": "queued"}


@app.post("/api/captures/url")
def import_url(payload: UrlCapture, background_tasks: BackgroundTasks) -> dict[str, Any]:
    content = f"# {payload.title}\n\nURL: {payload.url}\n\n{payload.context}".encode("utf-8")
    source_id, duplicate = create_source(f"{payload.title.strip()[:100]}.md", "url", content, url=payload.url)
    if duplicate:
        return {"id": source_id, "status": "duplicate"}
    job_id = str(uuid.uuid4())
    with db_connection() as db:
        db.execute("INSERT INTO processing_jobs(id, source_id, status, stage, created_at, updated_at) VALUES (?, ?, 'queued', 'waiting', ?, ?)", (job_id, source_id, now(), now()))
    background_tasks.add_task(process_source, source_id, job_id)
    return {"id": source_id, "job_id": job_id, "status": "queued"}


@app.get("/api/sources")
def list_sources() -> list[dict[str, Any]]:
    with db_connection() as db:
        rows = db.execute("SELECT * FROM sources ORDER BY created_at DESC").fetchall()
    return [source_payload(row) for row in rows]


@app.get("/api/jobs")
def list_jobs() -> list[dict[str, Any]]:
    with db_connection() as db:
        return [dict(row) for row in db.execute("SELECT * FROM processing_jobs ORDER BY created_at DESC LIMIT 100")]


@app.post("/api/sources/{source_id}/retry")
def retry_source(source_id: str, background_tasks: BackgroundTasks) -> dict[str, str]:
    with db_connection() as db:
        if not db.execute("SELECT 1 FROM sources WHERE id=?", (source_id,)).fetchone():
            raise HTTPException(404, "Source not found")
        job_id = str(uuid.uuid4())
        db.execute("UPDATE sources SET status='queued', error_message=NULL, updated_at=? WHERE id=?", (now(), source_id))
        db.execute("INSERT INTO processing_jobs(id, source_id, status, stage, created_at, updated_at) VALUES (?, ?, 'queued', 'waiting', ?, ?)", (job_id, source_id, now(), now()))
    background_tasks.add_task(process_source, source_id, job_id)
    return {"job_id": job_id, "status": "queued"}


@app.delete("/api/sources/{source_id}")
def delete_source(source_id: str) -> dict[str, bool]:
    with db_connection() as db:
        row = db.execute("SELECT managed_path FROM sources WHERE id=?", (source_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Source not found")
        db.execute("DELETE FROM chunks_fts WHERE chunk_id IN (SELECT id FROM chunks WHERE source_id=?)", (source_id,))
        db.execute("DELETE FROM sources WHERE id=?", (source_id,))
    Path(row["managed_path"]).unlink(missing_ok=True)
    return {"deleted": True}


@app.get("/api/sources/{source_id}/open")
def open_source(source_id: str):
    with db_connection() as db:
        row = db.execute("SELECT * FROM sources WHERE id=?", (source_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Source not found")
    return FileResponse(row["managed_path"], filename=row["filename"], media_type=row["mime_type"] or "application/octet-stream")


@app.get("/api/topics")
def list_topics() -> list[dict[str, Any]]:
    with db_connection() as db:
        rows = db.execute("""SELECT t.*, COUNT(st.source_id) source_count, n.status note_status
                              FROM topics t LEFT JOIN source_topics st ON st.topic_id=t.id
                              LEFT JOIN notes n ON n.topic_id=t.id GROUP BY t.id ORDER BY source_count DESC, t.name""").fetchall()
    return [dict(row) for row in rows]


@app.get("/api/cards")
def list_cards(q: str = "") -> list[dict[str, Any]]:
    with db_connection() as db:
        if q.strip():
            rows = db.execute(
                """SELECT c.*, s.filename, t.name AS topic_name FROM cards c
                   JOIN sources s ON s.id=c.source_id LEFT JOIN topics t ON t.id=c.topic_id
                   WHERE c.status='active' AND (c.title LIKE ? OR c.body LIKE ? OR s.filename LIKE ?)
                   ORDER BY c.updated_at DESC""",
                (f"%{q.strip()}%", f"%{q.strip()}%", f"%{q.strip()}%"),
            ).fetchall()
        else:
            rows = db.execute(
                """SELECT c.*, s.filename, t.name AS topic_name FROM cards c
                   JOIN sources s ON s.id=c.source_id LEFT JOIN topics t ON t.id=c.topic_id
                   WHERE c.status='active' ORDER BY c.created_at DESC"""
            ).fetchall()
    return [card_payload(row) for row in rows]


@app.get("/api/remember")
def remember_queue(limit: int = 5) -> list[dict[str, Any]]:
    limit = max(1, min(limit, 10))
    with db_connection() as db:
        rows = db.execute(
            """SELECT c.*, s.filename, t.name AS topic_name FROM cards c
               JOIN sources s ON s.id=c.source_id LEFT JOIN topics t ON t.id=c.topic_id
               WHERE c.status='active' AND c.next_review_at <= ?
               ORDER BY c.next_review_at ASC, c.created_at ASC LIMIT ?""",
            (now(), limit),
        ).fetchall()
    return [card_payload(row) for row in rows]


@app.get("/api/cards/{card_id}")
def get_card(card_id: str) -> dict[str, Any]:
    with db_connection() as db:
        row = db.execute(
            """SELECT c.*, s.filename, t.name AS topic_name FROM cards c
               JOIN sources s ON s.id=c.source_id LEFT JOIN topics t ON t.id=c.topic_id WHERE c.id=?""",
            (card_id,),
        ).fetchone()
    if not row:
        raise HTTPException(404, "Knowledge card not found")
    return card_payload(row)


@app.post("/api/cards/{card_id}/review")
def review_card(card_id: str) -> dict[str, Any]:
    with db_connection() as db:
        row = db.execute("SELECT review_count FROM cards WHERE id=? AND status='active'", (card_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Active knowledge card not found")
        count = row["review_count"] + 1
        due = review_timestamp(count)
        db.execute(
            "UPDATE cards SET review_count=?, last_seen_at=?, next_review_at=?, updated_at=? WHERE id=?",
            (count, now(), due, now(), card_id),
        )
    return {"reviewed": True, "review_count": count, "next_review_at": due}


@app.post("/api/cards/{card_id}/archive")
def archive_card(card_id: str) -> dict[str, bool]:
    with db_connection() as db:
        cursor = db.execute("UPDATE cards SET status='archived', updated_at=? WHERE id=?", (now(), card_id))
        if not cursor.rowcount:
            raise HTTPException(404, "Knowledge card not found")
    return {"archived": True}


@app.get("/api/topics/{topic_id}")
def get_topic(topic_id: str) -> dict[str, Any]:
    with db_connection() as db:
        topic = db.execute("SELECT * FROM topics WHERE id=?", (topic_id,)).fetchone()
        if not topic:
            raise HTTPException(404, "Topic not found")
        sources = db.execute("SELECT s.* FROM sources s JOIN source_topics st ON st.source_id=s.id WHERE st.topic_id=?", (topic_id,)).fetchall()
        draft = db.execute("""SELECT r.* FROM notes n JOIN note_revisions r ON r.note_id=n.id
                              WHERE n.topic_id=? AND r.status='draft' ORDER BY r.created_at DESC LIMIT 1""", (topic_id,)).fetchone()
        current = db.execute("""SELECT r.* FROM notes n JOIN note_revisions r ON r.id=n.current_revision_id
                                WHERE n.topic_id=?""", (topic_id,)).fetchone()
    return {"topic": dict(topic), "sources": [source_payload(row) for row in sources], "draft": dict(draft) if draft else None, "current_note": dict(current) if current else None}


@app.post("/api/topics/{topic_id}/note-drafts")
def generate_note_draft(topic_id: str) -> dict[str, Any]:
    try:
        return create_note_draft(topic_id)
    except KeyError:
        raise HTTPException(404, "Topic not found")
    except ValueError as error:
        raise HTTPException(422, str(error))


@app.post("/api/note-revisions/{revision_id}/approve")
def approve_note(revision_id: str) -> dict[str, bool]:
    with db_connection() as db:
        revision = db.execute("SELECT * FROM note_revisions WHERE id=?", (revision_id,)).fetchone()
        if not revision:
            raise HTTPException(404, "Note draft not found")
        db.execute("UPDATE note_revisions SET status='superseded' WHERE note_id=? AND status='approved'", (revision["note_id"],))
        db.execute("UPDATE note_revisions SET status='approved', reviewed_at=? WHERE id=?", (now(), revision_id))
        db.execute("UPDATE notes SET status='approved', current_revision_id=?, updated_at=? WHERE id=?", (revision_id, now(), revision["note_id"]))
        topic = db.execute("SELECT t.name FROM topics t JOIN notes n ON n.topic_id=t.id WHERE n.id=?", (revision["note_id"],)).fetchone()
    safe_name = re.sub(r"[^A-Za-z0-9 _-]", "", topic["name"]).strip() or "Untitled"
    (NOTES_DIR / f"{safe_name}.md").write_text(revision["body"], encoding="utf-8")
    return {"approved": True}


@app.post("/api/note-revisions/{revision_id}/reject")
def reject_note(revision_id: str) -> dict[str, bool]:
    with db_connection() as db:
        row = db.execute("SELECT note_id FROM note_revisions WHERE id=?", (revision_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Note draft not found")
        db.execute("UPDATE note_revisions SET status='rejected', reviewed_at=? WHERE id=?", (now(), revision_id))
    return {"rejected": True}


@app.get("/api/search")
def search(q: str = "") -> list[dict[str, Any]]:
    if not q.strip():
        return []
    chunks = query_rows(q, limit=6)
    return [
        {
            "chunk_id": row.get("chunk_id") or row.get("id"),
            "source_id": row["source_id"],
            "filename": row["filename"],
            "anchor": row.get("anchor") or "Source",
            "text": row["text"],
        }
        for row in chunks
    ]


@app.post("/api/questions")
def answer_question(payload: QuestionInput) -> dict[str, Any]:
    rows = query_rows(payload.question, limit=6)
    if not rows:
        answer = "I couldn't find supporting material in your Ren library yet. Import a relevant source or try different words."
        citations: list[dict[str, Any]] = []
    else:
        context, citations = build_grounded_context_and_citations(rows)
        ai_answer = provider_answer(payload.question, rows)
        answer = ai_answer or ("Based on your imported sources:\n\n" + "\n\n".join(r["text"][:650] for r in rows[:3]))
    with db_connection() as db:
        db.execute("INSERT INTO questions(id, question, answer, created_at) VALUES (?, ?, ?, ?)", (str(uuid.uuid4()), payload.question, answer, now()))
    return {"answer": answer, "citations": citations, "mode": "ai-grounded" if rows and api_key_available() and not answer.startswith("Based on") else "extractive"}


@app.get("/api/export")
def export_data() -> dict[str, Any]:
    with db_connection() as db:
        sources = [dict(row) for row in db.execute("SELECT id, filename, source_kind, original_path, url, status, created_at FROM sources")]
        topics = [dict(row) for row in db.execute("SELECT * FROM topics")]
        notes = [dict(row) for row in db.execute("SELECT * FROM note_revisions WHERE status='approved'")]
    return {"exported_at": now(), "sources": sources, "topics": topics, "approved_note_revisions": notes}

